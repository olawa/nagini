#!/usr/bin/env python3
"""Compare Nagini BCR-ABL1 caller results against Illumina MiSeq and PacBio Sequel.

This script takes:
1. Ground truth / validation table (e.g. table6_precision.txt) containing MiSeq and Sequel calls.
2. A Nagini analysis directory (or summary Excel workbook) run with all callers
   (Branford, VarDictRust, Rindels, bcftools).

It matches variants per sample across all callers and outputs a comparison table with:
- Allele frequencies (AF %) for MiSeq, Sequel, Branford, VarDictRust, Rindels, bcftools.
- Fusion isoforms (MiSeq, Sequel, Nagini).
- Compound phasing (Nagini CIS / TRANS).
- Export to both TSV and an Excel workbook with formatting.
"""

from __future__ import annotations

import argparse
import csv
import gzip
import os
import re
import sys
from pathlib import Path
from typing import Any, Dict, List, Optional, Set, Tuple


# Standard Branford amino acid 3-letter to 1-letter mapping
THREE_TO_ONE = {
    "Ala": "A", "Arg": "R", "Asn": "N", "Asp": "D", "Cys": "C",
    "Glu": "E", "Gln": "Q", "Gly": "G", "His": "H", "Ile": "I",
    "Leu": "L", "Lys": "K", "Met": "M", "Phe": "F", "Pro": "P",
    "Ser": "S", "Thr": "T", "Trp": "W", "Tyr": "Y", "Val": "V"
}


def normalize_aa(s: str) -> str:
    """Normalize amino acid change representation (e.g. p.Met351Thr -> M351T)."""
    if not s:
        return ""
    clean = str(s).strip()
    m = re.match(r"^p\.([A-Z][a-z]{2})(\d+)([A-Z][a-z]{2})$", clean)
    if m:
        ref, pos, alt = m.groups()
        return f"{THREE_TO_ONE.get(ref, ref)}{pos}{THREE_TO_ONE.get(alt, alt)}"
    m2 = re.match(r"^([A-Z])(\d+)([A-Z])$", clean)
    if m2:
        return clean
    return clean


def parse_vaf_percent(val: Any) -> Optional[float]:
    """Parse various VAF formats into a percentage float (0.0 to 100.0).

    Handles:
    - '96.97%' -> 96.97
    - '95,3'   -> 95.3
    - '0.946'  -> 94.6  (fraction <= 1.0 converted to percent)
    - 0.946    -> 94.6
    - 96.97    -> 96.97
    """
    if val is None:
        return None
    s = str(val).strip()
    if not s or s in ("-", "N/A", "n d", "nd", ".", "None", "ses ej"):
        return None
    is_explicit_percent = "%" in s
    s = s.replace("%", "").replace(",", ".").strip()
    try:
        num = float(s)
    except ValueError:
        return None

    if is_explicit_percent:
        return round(num, 2)
    # If it's a fraction between 0.0 and 1.0 (typical VCF AF), convert to percent
    if 0.0 < num <= 1.0:
        return round(num * 100.0, 2)
    return round(num, 2)


def load_branford_bed(bed_path: Path) -> Tuple[Dict[str, Tuple[str, int, str, str]], Dict[Tuple[str, int, str], str]]:
    """Load Branford BED file to map between AA change and genomic coordinates."""
    aa_to_coords: Dict[str, Tuple[str, int, str, str]] = {}
    coords_to_aa: Dict[Tuple[str, int, str], str] = {}

    if not bed_path.is_file():
        return aa_to_coords, coords_to_aa

    with open(bed_path, "r", encoding="utf-8") as f:
        for line in f:
            if line.startswith("#") or not line.strip():
                continue
            parts = line.strip().split("\t")
            if len(parts) >= 7:
                chrom = parts[0]
                pos = int(parts[2])  # 1-based end position
                ref = parts[4]
                alt = parts[5]
                aa_raw = parts[6]
                aa_short = normalize_aa(aa_raw)
                aa_to_coords[aa_short] = (chrom, pos, ref, alt)
                coords_to_aa[(chrom, pos, alt)] = aa_short
                coords_to_aa[(chrom.removeprefix("chr"), pos, alt)] = aa_short

    return aa_to_coords, coords_to_aa


def parse_table6(path: Path) -> List[Dict[str, Any]]:
    """Parse Table 6 (precision / clinical validation table with Sequel and MiSeq)."""
    with open(path, "r", encoding="utf-8") as f:
        lines = [l.strip() for l in f if l.strip()]

    start_idx = 0
    for idx, l in enumerate(lines):
        if l.startswith("Provnr") or "aa-variant" in l:
            start_idx = idx + 1

    records: List[Dict[str, Any]] = []

    def split_vars_and_afs(var_s: str, af_s: str) -> List[Tuple[str, Optional[float]]]:
        v_clean = var_s.strip()
        if not v_clean or v_clean.lower() in ("ingen variant", "ingen variant*", "ses ej", "låg cov", "snv i bcr"):
            return [(v_clean, None)]
        vars_list = [v.strip() for v in re.split(r"[/,\n]+", v_clean) if v.strip()]
        afs_list = [a.strip() for a in re.split(r"[/;\n\s]+", af_s) if a.strip()]
        res = []
        for i, v in enumerate(vars_list):
            af_val = parse_vaf_percent(afs_list[i]) if i < len(afs_list) else None
            res.append((normalize_aa(v), af_val))
        return res

    for l in lines[start_idx:]:
        parts = [p.strip() for p in l.split("\t")]
        if len(parts) < 8:
            continue
        provnr = parts[0]
        ratio = parts[1]
        seq_var_str = parts[2]
        seq_af_str = parts[3]
        seq_iso = parts[4].replace("-", "").strip()
        mis_var_str = parts[5]
        mis_af_str = parts[6]
        mis_iso = parts[7].replace("-", "").strip()

        seq_items = dict(split_vars_and_afs(seq_var_str, seq_af_str))
        mis_items = dict(split_vars_and_afs(mis_var_str, mis_af_str))

        all_vars = list(dict.fromkeys(list(seq_items.keys()) + list(mis_items.keys())))
        for v in all_vars:
            records.append({
                "provnr": provnr,
                "ratio": ratio,
                "variant": v,
                "sequel_af": seq_items.get(v),
                "sequel_iso": seq_iso,
                "miseq_af": mis_items.get(v),
                "miseq_iso": mis_iso,
            })

    return records


def parse_sample_mapping(run_dir: Path) -> Dict[str, Dict[str, str]]:
    """Parse samples.tsv and units.tsv to map sample ID to DNA Number and barcode."""
    mapping: Dict[str, Dict[str, str]] = {}

    samples_file = run_dir / "samples.tsv"
    if samples_file.is_file():
        with open(samples_file, "r", encoding="utf-8") as f:
            reader = csv.DictReader(f, delimiter="\t")
            for r in reader:
                s_id = r.get("sample", "").strip()
                dna = r.get("dna_number", r.get("patient_id", "")).strip()
                if s_id:
                    mapping[s_id] = {"dna_number": dna, "sample_id": s_id, "barcode": ""}

    units_file = run_dir / "units.tsv"
    if units_file.is_file():
        with open(units_file, "r", encoding="utf-8") as f:
            reader = csv.DictReader(f, delimiter="\t")
            for r in reader:
                s_id = r.get("sample", "").strip()
                bc = r.get("barcode", "").strip()
                if s_id in mapping:
                    mapping[s_id]["barcode"] = bc
                elif s_id:
                    mapping[s_id] = {"dna_number": "", "sample_id": s_id, "barcode": bc}

    return mapping


def parse_vcf_calls(vcf_path: Path, caller_name: str) -> List[Dict[str, Any]]:
    """Parse small variant VCF (VarDictRust, Rindels, bcftools)."""
    calls = []
    if not vcf_path.is_file():
        return calls

    open_fn = gzip.open if vcf_path.name.endswith(".gz") else open
    try:
        with open_fn(vcf_path, "rt", encoding="utf-8", errors="replace") as f:
            for line in f:
                if line.startswith("#"):
                    continue
                parts = line.strip().split("\t")
                if len(parts) < 8:
                    continue
                chrom, pos_s, _, ref, alt, _, filter_val, info_s = parts[:8]
                fmt_keys = parts[8].split(":") if len(parts) > 8 else []
                fmt_vals = parts[9].split(":") if len(parts) > 9 else []
                fmt_map = dict(zip(fmt_keys, fmt_vals))

                pos = int(pos_s)
                # Parse AF from format or info
                af = None
                if "AF" in fmt_map:
                    try:
                        af = parse_vaf_percent(fmt_map["AF"].split(",")[0])
                    except (ValueError, IndexError):
                        pass
                elif "VAF" in fmt_map:
                    af = parse_vaf_percent(fmt_map["VAF"].split(",")[0])
                elif "AD" in fmt_map and "DP" in fmt_map:
                    try:
                        dp = float(fmt_map["DP"])
                        ad = float(fmt_map["AD"].split(",")[1])
                        if dp > 0:
                            af = round((ad / dp) * 100.0, 2)
                    except (ValueError, IndexError):
                        pass

                calls.append({
                    "caller": caller_name,
                    "chrom": chrom,
                    "pos": pos,
                    "ref": ref,
                    "alt": alt,
                    "vaf": af,
                    "filter": filter_val,
                })
    except Exception as exc:
        sys.stderr.write(f"Warning: could not parse {vcf_path}: {exc}\n")

    return calls


def parse_branford_tsv(tsv_path: Path) -> List[Dict[str, Any]]:
    """Parse Branford caller TSV (all_sites.tsv or mutations.tsv)."""
    calls = []
    if not tsv_path.is_file():
        return calls
    with open(tsv_path, "r", encoding="utf-8") as f:
        reader = csv.DictReader(f, delimiter="\t")
        for r in reader:
            chrom = r.get("chrom", "")
            try:
                pos = int(r.get("pos", 0))
            except ValueError:
                continue
            ref = r.get("ref", "")
            alt = r.get("alt", "")
            aa = normalize_aa(r.get("aa_change", ""))
            vaf = parse_vaf_percent(r.get("vaf_percent", r.get("vaf", "")))
            status = r.get("status", "")
            calls.append({
                "caller": "Branford",
                "chrom": chrom,
                "pos": pos,
                "ref": ref,
                "alt": alt,
                "aa": aa,
                "vaf": vaf,
                "status": status,
            })
    return calls


def parse_nagini_summary_excel(xlsx_path: Path) -> Tuple[Dict[str, Dict[str, Any]], List[Dict[str, Any]]]:
    """Parse nagini_summary_report.xlsx to extract overview and PASS variants."""
    import openpyxl

    overview_map: Dict[str, Dict[str, Any]] = {}
    variants: List[Dict[str, Any]] = []

    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    if "Overview" in wb.sheetnames:
        sheet = wb["Overview"]
        header = [cell.value for cell in sheet[1]]
        dna_idx = header.index("DNA Number") if "DNA Number" in header else None
        s_idx = header.index("Sample ID") if "Sample ID" in header else None
        bc_idx = header.index("Barcode") if "Barcode" in header else None
        iso_idx = header.index("Primary Isoform") if "Primary Isoform" in header else None
        phase_idx = header.index("Compound Phasing") if "Compound Phasing" in header else None
        status_idx = header.index("Status") if "Status" in header else None

        for row in sheet.iter_rows(min_row=2, values_only=True):
            if not row or not any(row):
                continue
            dna = str(row[dna_idx]).strip() if dna_idx is not None and row[dna_idx] else ""
            sid = str(row[s_idx]).strip() if s_idx is not None and row[s_idx] else ""
            bc = str(row[bc_idx]).strip() if bc_idx is not None and row[bc_idx] else ""
            iso = str(row[iso_idx]).strip() if iso_idx is not None and row[iso_idx] else ""
            phase = str(row[phase_idx]).strip() if phase_idx is not None and row[phase_idx] else ""
            stat = str(row[status_idx]).strip() if status_idx is not None and row[status_idx] else ""

            entry = {
                "dna_number": dna,
                "sample_id": sid,
                "barcode": bc,
                "isoform": iso,
                "phasing": phase,
                "status": stat,
            }
            if dna:
                overview_map[dna] = entry
            if sid:
                overview_map[sid] = entry

    if "PASS Variants" in wb.sheetnames:
        sheet = wb["PASS Variants"]
        header = [cell.value for cell in sheet[1]]
        dna_idx = header.index("DNA Number") if "DNA Number" in header else None
        s_idx = header.index("Sample ID") if "Sample ID" in header else None
        caller_idx = header.index("Caller") if "Caller" in header else None
        gene_idx = header.index("Gene") if "Gene" in header else None
        chr_idx = header.index("Chr") if "Chr" in header else None
        pos_idx = header.index("Position") if "Position" in header else None
        ref_idx = header.index("Ref") if "Ref" in header else None
        alt_idx = header.index("Alt") if "Alt" in header else None
        vaf_idx = header.index("VAF") if "VAF" in header else None
        aa_idx = header.index("Consequence / AA Change") if "Consequence / AA Change" in header else None

        for row in sheet.iter_rows(min_row=2, values_only=True):
            if not row or not any(row):
                continue
            dna = str(row[dna_idx]).strip() if dna_idx is not None and row[dna_idx] else ""
            sid = str(row[s_idx]).strip() if s_idx is not None and row[s_idx] else ""
            caller = str(row[caller_idx]).strip() if caller_idx is not None and row[caller_idx] else ""
            chrom = str(row[chr_idx]).strip() if chr_idx is not None and row[chr_idx] else ""
            try:
                pos = int(row[pos_idx]) if pos_idx is not None and row[pos_idx] else 0
            except ValueError:
                pos = 0
            ref = str(row[ref_idx]).strip() if ref_idx is not None and row[ref_idx] else ""
            alt = str(row[alt_idx]).strip() if alt_idx is not None and row[alt_idx] else ""
            vaf = parse_vaf_percent(row[vaf_idx]) if vaf_idx is not None else None
            aa = normalize_aa(row[aa_idx]) if aa_idx is not None and row[aa_idx] else ""

            variants.append({
                "dna_number": dna,
                "sample_id": sid,
                "caller": caller,
                "chrom": chrom,
                "pos": pos,
                "ref": ref,
                "alt": alt,
                "vaf": vaf,
                "aa": aa,
            })

    return overview_map, variants


def build_comparison(
    truth_records: List[Dict[str, Any]],
    run_dir: Path,
    summary_xlsx: Optional[Path] = None,
    aa_to_coords: Optional[Dict[str, Tuple[str, int, str, str]]] = None,
    coords_to_aa: Optional[Dict[Tuple[str, int, str], str]] = None,
) -> List[Dict[str, Any]]:
    """Assemble the unified comparison table across MiSeq, Sequel, and Nagini callers."""
    aa_to_coords = aa_to_coords or {}
    coords_to_aa = coords_to_aa or {}

    # Sample mapping from run_dir
    sample_mapping = parse_sample_mapping(run_dir)
    dna_to_sid: Dict[str, str] = {}
    sid_to_dna: Dict[str, str] = {}
    for sid, info in sample_mapping.items():
        dna = info.get("dna_number", "")
        if dna:
            dna_to_sid[dna] = sid
            sid_to_dna[sid] = dna

    # Look for summary xlsx if not explicitly provided
    if not summary_xlsx:
        cand_dirs = [run_dir, run_dir / "reports", run_dir / "results" / "reports"]
        for cd in cand_dirs:
            if not cd.is_dir():
                continue
            for cand_name in ("nagini_summary_report.xlsx",):
                cand = cd / cand_name
                if cand.is_file():
                    summary_xlsx = cand
                    break
            if summary_xlsx:
                break
            for p in sorted(cd.glob("*_summary_report.xlsx")):
                if p.is_file():
                    summary_xlsx = p
                    break
            if summary_xlsx:
                break

    overview_map: Dict[str, Dict[str, Any]] = {}
    excel_variants: List[Dict[str, Any]] = []
    if summary_xlsx and summary_xlsx.is_file():
        try:
            overview_map, excel_variants = parse_nagini_summary_excel(summary_xlsx)
            for k, info in overview_map.items():
                dna = info.get("dna_number")
                sid = info.get("sample_id")
                if dna and sid:
                    dna_to_sid[dna] = sid
                    sid_to_dna[sid] = dna
        except Exception as exc:
            sys.stderr.write(f"Warning: could not read {summary_xlsx}: {exc}\n")

    # Group comparison by (provnr, variant)
    comparison_rows: List[Dict[str, Any]] = []

    for tr in truth_records:
        provnr = tr["provnr"]
        ratio = tr["ratio"]
        var_name = tr["variant"]
        miseq_af = tr["miseq_af"]
        miseq_iso = tr["miseq_iso"]
        sequel_af = tr["sequel_af"]
        sequel_iso = tr["sequel_iso"]

        # Resolve sample ID and overview info
        sid = dna_to_sid.get(provnr, "")
        if not sid and provnr in overview_map:
            sid = overview_map[provnr].get("sample_id", "")
        if not sid:
            # Fallback check if provnr is sid
            sid = provnr if provnr in sample_mapping else ""

        ov = overview_map.get(provnr) or overview_map.get(sid) or {}
        barcode = ov.get("barcode") or sample_mapping.get(sid, {}).get("barcode", "")
        nagini_iso = ov.get("isoform", "")
        nagini_phase = ov.get("phasing", "")

        # Fallback to direct TSV files if not already populated from Excel overview
        sample_candidates = [f"{sid}_T", sid] if sid else []
        if not nagini_iso:
            for sp in sample_candidates:
                iso_path = run_dir / "reports" / "bcr_abl" / f"{sp}.fusion_isoforms.tsv"
                if not iso_path.is_file():
                    iso_path = run_dir / "results" / "reports" / "bcr_abl" / f"{sp}.fusion_isoforms.tsv"
                if iso_path.is_file():
                    try:
                        with open(iso_path, "r", encoding="utf-8") as isof:
                            reader = csv.DictReader(isof, delimiter="\t")
                            for ir in reader:
                                if ir.get("status") == "PRIMARY":
                                    nagini_iso = ir.get("fusion_isoform", "")
                                    break
                    except Exception:
                        pass
                if nagini_iso:
                    break

        if not nagini_phase:
            for sp in sample_candidates:
                ph_path = run_dir / "reports" / "bcr_abl" / f"{sp}.compound_phasing.tsv"
                if not ph_path.is_file():
                    ph_path = run_dir / "results" / "reports" / "bcr_abl" / f"{sp}.compound_phasing.tsv"
                if ph_path.is_file():
                    try:
                        with open(ph_path, "r", encoding="utf-8") as phf:
                            reader = csv.DictReader(phf, delimiter="\t")
                            cis_pairs = []
                            trans_pairs = []
                            for pr in reader:
                                state = pr.get("phase_state", "")
                                m1 = pr.get("mutation_1", "")
                                m2 = pr.get("mutation_2", "")
                                if "CIS" in state.upper():
                                    cis_pairs.append(f"{m1}+{m2}")
                                elif "TRANS" in state.upper():
                                    trans_pairs.append(f"{m1}+{m2}")
                            if cis_pairs:
                                nagini_phase = f"CIS ({', '.join(cis_pairs)})"
                            elif trans_pairs:
                                nagini_phase = f"TRANS ({', '.join(trans_pairs)})"
                    except Exception:
                        pass
                if nagini_phase:
                    break

        # Look up expected coordinates for this variant
        norm_v = normalize_aa(var_name)
        coord_tuple = aa_to_coords.get(norm_v)
        chrom_exp = coord_tuple[0] if coord_tuple else ""
        pos_exp = coord_tuple[1] if coord_tuple else None
        ref_exp = coord_tuple[2] if coord_tuple else ""
        alt_exp = coord_tuple[3] if coord_tuple else ""

        # Initialize caller AFs
        branford_af = None
        vardict_af = None
        rindels_af = None
        bcftools_af = None

        # 1. Check Excel PASS variants
        for ev in excel_variants:
            if ev.get("dna_number") == provnr or ev.get("sample_id") == sid:
                c_caller = ev.get("caller", "")
                c_aa = normalize_aa(ev.get("aa", ""))
                c_pos = ev.get("pos")
                is_match = False
                if norm_v and c_aa and norm_v == c_aa:
                    is_match = True
                elif pos_exp and c_pos == pos_exp:
                    is_match = True

                if is_match:
                    if "branford" in c_caller.lower():
                        branford_af = ev.get("vaf")
                    elif "vardict" in c_caller.lower():
                        vardict_af = ev.get("vaf")
                    elif "rindels" in c_caller.lower():
                        rindels_af = ev.get("vaf")
                    elif "bcftools" in c_caller.lower():
                        bcftools_af = ev.get("vaf")

        # 2. Check direct files if caller AF is still missing
        sample_pref = f"{sid}_T" if sid else ""
        if sample_pref:
            # Branford all_sites / mutations
            if branford_af is None:
                b_mut_path = run_dir / "reports" / "bcr_abl" / f"{sample_pref}.branford_mutations.tsv"
                b_calls = parse_branford_tsv(b_mut_path)
                for bc in b_calls:
                    if (norm_v and bc.get("aa") == norm_v) or (pos_exp and bc.get("pos") == pos_exp):
                        branford_af = bc.get("vaf")
                        break
            if branford_af is None and pos_exp:
                b_all_path = run_dir / "reports" / "bcr_abl" / f"{sample_pref}.branford_all_sites.tsv"
                b_calls = parse_branford_tsv(b_all_path)
                for bc in b_calls:
                    if bc.get("pos") == pos_exp and bc.get("alt") == alt_exp:
                        branford_af = bc.get("vaf")
                        break

            # VarDictRust VCF
            if vardict_af is None:
                vd_path = run_dir / "snv_indels" / "vardictrust" / f"{sample_pref}.bcr_abl.vcf.gz"
                if not vd_path.is_file():
                    vd_path = run_dir / "snv_indels" / "vardictrust" / f"{sample_pref}.bcr_abl.vcf"
                vd_calls = parse_vcf_calls(vd_path, "VarDictRust")
                for vc in vd_calls:
                    if pos_exp and vc.get("pos") == pos_exp and vc.get("alt") == alt_exp:
                        vardict_af = vc.get("vaf")
                        break

            # Rindels VCF
            if rindels_af is None:
                r_path = run_dir / "reports" / "bcr_abl" / f"{sample_pref}.rindels.vcf.gz"
                if not r_path.is_file():
                    r_path = run_dir / "reports" / "bcr_abl" / f"{sample_pref}.rindels.vcf"
                if not r_path.is_file():
                    r_path = run_dir / "cnv" / "rindels" / f"{sample_pref}.vcf.gz"
                r_calls = parse_vcf_calls(r_path, "Rindels")
                for rc in r_calls:
                    if pos_exp and rc.get("pos") == pos_exp and rc.get("alt") == alt_exp:
                        rindels_af = rc.get("vaf")
                        break

            # bcftools pileup
            if bcftools_af is None:
                bcf_path = run_dir / "reports" / "bcr_abl" / f"{sample_pref}.bcftools_pileup.tsv"
                if bcf_path.is_file():
                    with open(bcf_path, "r", encoding="utf-8") as bf:
                        reader = csv.DictReader(bf, delimiter="\t")
                        for br in reader:
                            try:
                                b_pos = int(br.get("pos", 0))
                            except ValueError:
                                continue
                            if pos_exp and b_pos == pos_exp:
                                bcftools_af = parse_vaf_percent(br.get("vaf", br.get("af", "")))
                                break

        coord_str = f"{chrom_exp}:{pos_exp} {ref_exp}>{alt_exp}" if pos_exp else "-"

        comparison_rows.append({
            "dna_number": provnr,
            "sample_id": sid,
            "barcode": barcode,
            "ratio": ratio,
            "variant": var_name,
            "genomic_pos": coord_str,
            "miseq_af": miseq_af,
            "sequel_af": sequel_af,
            "branford_af": branford_af,
            "vardict_af": vardict_af,
            "rindels_af": rindels_af,
            "bcftools_af": bcftools_af,
            "miseq_iso": miseq_iso,
            "sequel_iso": sequel_iso,
            "nagini_iso": nagini_iso,
            "nagini_phasing": nagini_phase,
        })

    return comparison_rows


def format_af(val: Optional[float]) -> str:
    """Format AF float as string percentage."""
    if val is None:
        return "-"
    return f"{val:.2f}%"


def write_tsv(rows: List[Dict[str, Any]], out_path: Path) -> None:
    """Write comparison rows to TSV."""
    fieldnames = [
        "DNA Number", "Sample ID", "Barcode", "Ratio / IS", "Variant", "Genomic Position",
        "MiSeq AF (%)", "Sequel AF (%)", "Branford AF (%)", "VarDictRust AF (%)",
        "Rindels AF (%)", "bcftools AF (%)",
        "MiSeq Isoform", "Sequel Isoform", "Nagini Isoform", "Nagini Compound Phasing"
    ]
    with open(out_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f, delimiter="\t")
        writer.writerow(fieldnames)
        for r in rows:
            writer.writerow([
                r["dna_number"],
                r["sample_id"],
                r["barcode"],
                r["ratio"],
                r["variant"],
                r["genomic_pos"],
                format_af(r["miseq_af"]),
                format_af(r["sequel_af"]),
                format_af(r["branford_af"]),
                format_af(r["vardict_af"]),
                format_af(r["rindels_af"]),
                format_af(r["bcftools_af"]),
                r["miseq_iso"],
                r["sequel_iso"],
                r["nagini_iso"],
                r["nagini_phasing"],
            ])


def write_excel(rows: List[Dict[str, Any]], out_path: Path) -> None:
    """Write comparison rows to a styled Excel workbook."""
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "BCR-ABL1 Caller Comparison"

    headers = [
        "DNA Number", "Sample ID", "Barcode", "Ratio / IS", "Variant", "Genomic Position",
        "MiSeq AF (%)", "Sequel AF (%)", "Branford AF (%)", "VarDictRust AF (%)",
        "Rindels AF (%)", "bcftools AF (%)",
        "MiSeq Isoform", "Sequel Isoform", "Nagini Isoform", "Nagini Compound Phasing"
    ]

    # Styling
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    zebra_fill = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")
    regular_font = Font(name="Calibri", size=10)
    bold_font = Font(name="Calibri", size=10, bold=True)
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    right_align = Alignment(horizontal="right", vertical="center")

    thin_border = Border(
        left=Side(style="thin", color="E0E0E0"),
        right=Side(style="thin", color="E0E0E0"),
        top=Side(style="thin", color="E0E0E0"),
        bottom=Side(style="thin", color="E0E0E0"),
    )

    ws.append(headers)
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align

    for r_idx, r in enumerate(rows, start=2):
        row_vals = [
            r["dna_number"],
            r["sample_id"],
            r["barcode"],
            r["ratio"],
            r["variant"],
            r["genomic_pos"],
            format_af(r["miseq_af"]),
            format_af(r["sequel_af"]),
            format_af(r["branford_af"]),
            format_af(r["vardict_af"]),
            format_af(r["rindels_af"]),
            format_af(r["bcftools_af"]),
            r["miseq_iso"],
            r["sequel_iso"],
            r["nagini_iso"],
            r["nagini_phasing"],
        ]
        ws.append(row_vals)
        use_fill = zebra_fill if (r_idx % 2 == 0) else None

        for c_idx, val in enumerate(row_vals, start=1):
            cell = ws.cell(row=r_idx, column=c_idx)
            cell.font = bold_font if c_idx in (1, 5) else regular_font
            cell.border = thin_border
            if use_fill:
                cell.fill = use_fill
            if c_idx in (1, 2, 3, 4, 13, 14, 15, 16):
                cell.alignment = center_align
            elif c_idx in (7, 8, 9, 10, 11, 12):
                cell.alignment = right_align
            else:
                cell.alignment = left_align

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    ws.freeze_panes = "A2"
    wb.save(out_path)


def print_markdown_summary(rows: List[Dict[str, Any]]) -> None:
    """Print Markdown table to stdout."""
    headers = [
        "DNA Number", "Sample ID", "Ratio", "Variant",
        "MiSeq AF", "Sequel AF", "Branford AF", "VarDict AF", "Rindels AF", "bcftools AF",
        "MiSeq Iso", "Nagini Iso", "Nagini Phasing"
    ]
    print("\n" + " | ".join(headers))
    print("| " + " | ".join(["---"] * len(headers)) + " |")
    for r in rows:
        line = [
            r["dna_number"],
            r["sample_id"] or "-",
            r["ratio"],
            r["variant"],
            format_af(r["miseq_af"]),
            format_af(r["sequel_af"]),
            format_af(r["branford_af"]),
            format_af(r["vardict_af"]),
            format_af(r["rindels_af"]),
            format_af(r["bcftools_af"]),
            r["miseq_iso"],
            r["nagini_iso"] or "-",
            r["nagini_phasing"] or "-",
        ]
        print(" | ".join(line))
    print()


def main(argv: Optional[List[str]] = None) -> int:
    parser = argparse.ArgumentParser(
        description="Compare Nagini BCR-ABL1 calls across callers against MiSeq and Sequel."
    )
    parser.add_argument(
        "--run-dir", "-r", type=Path, default=Path("."),
        help="Path to Nagini run directory (default: current directory)"
    )
    parser.add_argument(
        "--truth-table", "-t", type=Path, default=None,
        help="Path to Table 6 text file (default: searches ~/data/bcr-abl/table6_precision.txt or local)"
    )
    parser.add_argument(
        "--summary-xlsx", "-s", type=Path, default=None,
        help="Path to Nagini summary report workbook (optional, auto-detected from run-dir if omitted)"
    )
    parser.add_argument(
        "--bed", "-b", type=Path, default=None,
        help="Path to Branford positions BED file (default: workflow/resources/bcr_abl/branford_positions_hg19.bed)"
    )
    parser.add_argument(
        "--output-tsv", "-o", type=Path, default=Path("bcr_abl_caller_comparison.tsv"),
        help="Output TSV path (default: bcr_abl_caller_comparison.tsv)"
    )
    parser.add_argument(
        "--output-xlsx", "-x", type=Path, default=Path("bcr_abl_caller_comparison.xlsx"),
        help="Output Excel workbook path (default: bcr_abl_caller_comparison.xlsx)"
    )

    args = parser.parse_args(argv)

    # Resolve truth table path
    truth_path = args.truth_table
    if not truth_path or not truth_path.is_file():
        candidates = [
            Path("~/data/bcr-abl/table6_precision.txt").expanduser(),
            Path("table6_precision.txt"),
            Path("../table6_precision.txt"),
        ]
        for c in candidates:
            if c.is_file():
                truth_path = c
                break

    if not truth_path or not truth_path.is_file():
        sys.stderr.write("Error: could not find truth table file. Specify with --truth-table /path/to/table6_precision.txt\n")
        return 1

    # Resolve Branford BED
    bed_path = args.bed
    if not bed_path or not bed_path.is_file():
        repo_root = Path(__file__).resolve().parent.parent
        bed_candidate = repo_root / "workflow" / "resources" / "bcr_abl" / "branford_positions_hg19.bed"
        if bed_candidate.is_file():
            bed_path = bed_candidate

    aa_to_coords = {}
    coords_to_aa = {}
    if bed_path and bed_path.is_file():
        aa_to_coords, coords_to_aa = load_branford_bed(bed_path)

    # Parse truth table
    truth_records = parse_table6(truth_path)
    if not truth_records:
        sys.stderr.write(f"Error: no records parsed from truth table {truth_path}\n")
        return 1

    # Build comparison
    comp_rows = build_comparison(
        truth_records=truth_records,
        run_dir=args.run_dir,
        summary_xlsx=args.summary_xlsx,
        aa_to_coords=aa_to_coords,
        coords_to_aa=coords_to_aa,
    )

    # Outputs
    print_markdown_summary(comp_rows)

    write_tsv(comp_rows, args.output_tsv)
    print(f"[Nagini] Saved TSV comparison table to: {args.output_tsv}")

    try:
        write_excel(comp_rows, args.output_xlsx)
        print(f"[Nagini] Saved Excel comparison table to: {args.output_xlsx}")
    except Exception as exc:
        sys.stderr.write(f"Warning: could not write Excel file: {exc}\n")

    return 0


if __name__ == "__main__":
    sys.exit(main())
